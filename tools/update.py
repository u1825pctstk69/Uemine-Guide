"""配置図Excelからガイドページを更新する

使い方:
  python3 tools/update.py 配置図.xlsx   … 配置図から作り直す
  python3 tools/update.py               … 前回のデータ(data/data.json)のまま、週間おすすめ台(data/weekly.json)だけ反映

1. Excelの中から「B1が入替日・D1の日付が最も新しいシート」を選ぶ
2. tools/build.js（= gas/Code.gs と同じ解析）で機種データを作る
3. 機種マスタ(data/master.csv)に無い型式があれば一覧を出して止まる
   → 型式ごとにスペック／タイプとP-WORLD機種IDを調べて master.csv に追記してから再実行
4. index.html と smasuro20/index.html を書き出す（git commit / push は別途）
"""
import sys, os, re, json, datetime, subprocess, tempfile
import openpyxl

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))


def pick_sheet(wb):
    best, best_d = None, None
    for ws in wb.worksheets:
        if str(ws['B1'].value or '').strip() != '入替日':
            continue
        v = ws['D1'].value
        d = v if isinstance(v, datetime.datetime) else (datetime.datetime(1899, 12, 30) + datetime.timedelta(days=v) if isinstance(v, (int, float)) else None)
        if d and (best_d is None or d > best_d):
            best, best_d = ws, d
    return best or wb.worksheets[0]


def dump_grid(ws):
    grid = []
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=ws.max_column):
        out = []
        for c in row:
            v = c.value
            if isinstance(v, datetime.datetime):
                v = {'__date': v.isoformat()}
            elif v is None:
                v = ''
            out.append(v)
        grid.append(out)
    return grid


def wrap(template, data_json, api, theme, weekly):
    body = template.replace('/*__DATA__*/', data_json).replace('/*__API__*/', api).replace('/*__WEEKLY__*/null', weekly)
    title = re.search(r'<title>.*?</title>', body).group(0)
    body = body.replace(title, '', 1)
    links = re.findall(r'<link[^>]+>\n?', body)
    for l in links:
        body = body.replace(l, '', 1)
    style = re.search(r'<style>.*?</style>', body, re.S).group(0)
    body = body.replace(style, '', 1)
    return f'''<!DOCTYPE html>
<html lang="ja">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width,initial-scale=1,viewport-fit=cover">
<meta name="theme-color" content="{theme}">
{title}
{''.join(links)}{style}
<style>[hidden]{{display:none!important}}img{{max-width:100%}}:root{{padding-top:env(safe-area-inset-top,0px);padding-bottom:env(safe-area-inset-bottom,0px)}}</style>
</head>
<body>
{body.strip()}
</body>
</html>
'''


def main():
    data_file = os.path.join(ROOT, 'data', 'data.json')
    if len(sys.argv) > 1:
        wb = openpyxl.load_workbook(sys.argv[1], data_only=True)
        ws = pick_sheet(wb)
        print('対象シート:', ws.title)
        tmp = tempfile.mkdtemp()
        gfile, ofile = os.path.join(tmp, 'grid.json'), os.path.join(tmp, 'data.json')
        json.dump(dump_grid(ws), open(gfile, 'w'), ensure_ascii=False)
        r = subprocess.run(['node', os.path.join(ROOT, 'tools', 'build.js'), gfile,
                            os.path.join(ROOT, 'data', 'master.csv'), ofile, ws.title])
        if r.returncode == 2:
            print('\n→ 未登録の型式を data/master.csv に追記してから、もう一度実行してください。')
            sys.exit(2)
        if r.returncode != 0:
            sys.exit(r.returncode)
        open(data_file, 'w').write(open(ofile).read())
    data = open(data_file).read()
    api_file = os.path.join(ROOT, 'data', 'api_url.txt')
    api = open(api_file).read().strip() if os.path.exists(api_file) else ''
    wk_file = os.path.join(ROOT, 'data', 'weekly.json')
    weekly = json.dumps(json.load(open(wk_file)), ensure_ascii=False) if os.path.exists(wk_file) else 'null'
    if weekly != 'null':
        names = {m['name'] for m in json.loads(data)['machines']}
        for it in json.loads(weekly).get('items', []):
            if it['name'] not in names:
                print('注意: 週間おすすめの機種が設置機種に見つかりません →', it['name'])
    tdir = os.path.join(ROOT, 'tools', 'templates')
    open(os.path.join(ROOT, 'index.html'), 'w').write(wrap(open(os.path.join(tdir, 'guide.html')).read(), data, api, '#B01030', weekly))
    os.makedirs(os.path.join(ROOT, 'smasuro20'), exist_ok=True)
    open(os.path.join(ROOT, 'smasuro20', 'index.html'), 'w').write(wrap(open(os.path.join(tdir, 'smasuro20.html')).read(), data, api, '#2F5BD3', weekly))
    print('index.html / smasuro20/index.html を更新しました。')
    # お試し: 台のQR用ページ（p5/）。テンプレートがあるときだけ作る
    p5 = os.path.join(tdir, 'p5.html')
    if os.path.exists(p5):
        vf = os.path.join(ROOT, 'data', 'vote.json')
        vote = json.dumps(json.load(open(vf)), ensure_ascii=False) if os.path.exists(vf) else 'null'
        lf = os.path.join(ROOT, 'data', 'log_url.txt')
        logu = open(lf).read().strip() if os.path.exists(lf) else ''
        t = open(p5).read().replace('/*__VOTE__*/null', vote).replace('/*__LOG__*/', logu)
        os.makedirs(os.path.join(ROOT, 'p5'), exist_ok=True)
        open(os.path.join(ROOT, 'p5', 'index.html'), 'w').write(wrap(t, data, api, '#0F9D8A', weekly))
        print('p5/index.html（お試し・台のQR用）を更新しました。')


if __name__ == '__main__':
    main()
